#!/usr/bin/env python3
"""
One-off, repeatable converter: the two curated Coding Hours workbooks -> the JSON files
this directory commits (group1.json, group2.json).

The JSON is the version-controlled source of truth the importer reads; the workbooks are
not needed at deploy time. Nothing is corrected or normalised here — every value is copied
as supplied, and the importer's validator reports (never fixes) anything wrong.

  python3 xlsx_to_json.py <Kalvium_Dojo_LeetCode_Daily_Question_Sheet.xlsx> <Group2_Daily_Question_Sheet.xlsx>

Requires: openpyxl (conversion only — not a project dependency).
"""
import json, os, re, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SLUG = re.compile(r"leetcode\.com/problems/([^/?#]+)")

def slug_of(url):
    m = SLUG.search(url or "")
    return m.group(1) if m else None

def group1(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Daily Questions"]
    rows = list(ws.iter_rows(values_only=True))
    h = {name: i for i, name in enumerate(rows[0])}
    out = []
    for n, r in enumerate(rows[1:], start=2):
        out.append({
            "sourceRow": n,
            "leetcodeNumber": r[h["LeetCode #"]], "title": r[h["Problem"]],
            "difficulty": r[h["Difficulty"]], "topic": r[h["Topic"]], "pattern": r[h["Pattern"]],
            "dayFocus": r[h["Day Focus"]], "url": r[h["URL"]], "titleSlug": slug_of(r[h["URL"]]),
            "belt": r[h["Belt"]], "week": r[h["Week"]], "day": r[h["Day"]],
            "position": r[h["Q#"]], "usage": r[h["Usage"]],
        })
    return out

def group2(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Daily Plan"]   # not read-only: hyperlinks
    head = [c.value for c in ws[1]]
    h = {name: i for i, name in enumerate(head)}
    out = []
    for n, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if row[0].value is None:
            continue
        link = row[h["Problem"]].hyperlink.target if row[h["Problem"]].hyperlink else None
        if link is None:  # the hyperlink may sit on the LC # cell instead
            link = row[h["LC #"]].hyperlink.target if row[h["LC #"]].hyperlink else None
        v = lambda k: row[h[k]].value
        out.append({
            "sourceRow": n, "week": v("Week"), "weekday": v("Day"), "dayNumber": v("Day #"),
            "dailyTheme": v("Daily theme"), "position": v("Q#"), "role": v("Role"),
            "leetcodeNumber": v("LC #"), "title": v("Problem"), "difficulty": v("Difficulty"),
            "url": link, "titleSlug": slug_of(link),
        })
    return out

if __name__ == "__main__":
    g1, g2 = group1(sys.argv[1]), group2(sys.argv[2])
    for name, data in (("group1.json", g1), ("group2.json", g2)):
        with open(os.path.join(HERE, name), "w") as fh:
            json.dump(data, fh, indent=1, ensure_ascii=False); fh.write("\n")
        print(name, len(data), "rows")

"""
Workbook writer for build_five_day.py — the reviewable backup of the five-day Group 1 plan.
Imported by that script; needs openpyxl (conversion/report only, not a project dependency).
"""
import collections
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BELT_NAMES = {1: "Yellow Belt", 2: "Orange Belt", 3: "Green Belt", 4: "Purple Belt", 5: "Blue Belt",
              6: "Brown Belt", 7: "Sr. Brown Belt", 8: "Red Belt", 9: "Sr. Red Belt", 10: "Black Belt"}

# pattern substring -> concept-table topic, per belt. First match wins. Numbers are the table's own.
RULES = {
 1: [("Functions", "5. Functions & Parameter Passing"), ("Conditionals", "4. Conditionals"),
     ("Arithmetic", "2. Operators & Expressions (and 1. Variables & Data Types)")],
 2: [("Pattern Problems", "2. Pattern Printing"), ("Complexity", "5. Complexity Basics"),
     ("Loops", "1. Loops & Control Flow"), ("Sorting", "3. Elementary Sorting"),
     ("GCD", "4. Basic Math & Number Theory — GCD / LCM"), ("Primes", "4. Basic Math & Number Theory — prime check"),
     ("Sieve", "4. Basic Math & Number Theory — Sieve of Eratosthenes"),
     ("Digit", "4. Basic Math & Number Theory — digit problems"),
     ("Modular", "4. Basic Math & Number Theory — modular arithmetic"),
     ("Fast Exponentiation", "4. Basic Math & Number Theory — fast exponentiation")],
 3: [("Frequency", "4. Frequency Counting"), ("Brute Force vs", "5. Complexity Growth"),
     ("Matri", "3. Matrices / 2D Arrays"), ("Spiral", "3. Matrices / 2D Arrays"),
     ("Palindrome", "2. Strings"), ("String", "2. Strings"), ("", "1. Arrays")],
 4: [("Subarrays (Brute", "1. Subarrays & Subsequences"), ("Prefix Sum + Hashing", "3. Hashing — subarray with a given sum"),
     ("Prefix", "2. Prefix & Suffix Sums"), ("Hashing", "3. Hashing"),
     ("Kadane", "4. Classic Array Algorithms — Kadane's"), ("Dutch", "4. Classic Array Algorithms — Dutch National Flag"),
     ("Moore", "4. Classic Array Algorithms — Moore's Voting"), ("Counting Sort", "8. Counting Sort & Radix Sort"),
     ("Two Pointers", "5. Two-Pointer Technique"), ("Sliding Window", "6. Sliding Window"), ("Binary Search", "7. Binary Search")],
 5: [("Hare-Tortoise", "5. Hare-Tortoise Method"), ("Fast & Slow", "5. Hare-Tortoise Method (middle / n-th from end)"),
     ("Monotonic Stack", "4. Monotonic Stack"), ("Monotonic Deque", "3. Queues & Deque — sliding window maximum"),
     ("Queue", "3. Queues & Deque"), ("Stack", "2. Stacks"), ("", "1. Linked Lists")],
 6: [("Recursion (Base", "1. Recursion Basics"), ("Recursion Tree", "1. Recursion Basics (recursion tree)"),
     ("Recursion (Reduction", "2. Core Recursive Problems"), ("Recursion (Strings", "2. Core Recursive Problems"),
     ("Recursive Parsing", "1./2. Recursion (nested structures)"), ("Recursion (Challenge", "1./2. Recursion (challenge)"),
     ("Merge Sort", "3. Divide & Conquer — Merge Sort"), ("Merging Sorted", "3. Divide & Conquer — merge step"),
     ("Divide", "3. Divide & Conquer"), ("Quickselect", "3. Divide & Conquer — Quickselect"),
     ("Partitioning (Sort", "3. Divide & Conquer — partitioning"),
     ("Backtracking (Subsets", "4. Backtracking: Generating Possibilities"), ("Backtracking (Permutations", "4. Backtracking: Generating Possibilities"),
     ("Backtracking (Combinations", "4. Backtracking: Generating Possibilities"),
     ("Backtracking (Grid", "5. Backtracking: Searching with Rules"), ("Backtracking (Partitioning", "5. Backtracking: Searching with Rules"),
     ("Backtracking (String Partitioning", "5. Backtracking: Searching with Rules"),
     ("Pruning", "6. Backtracking Pattern (pruning)"), ("Backtracking (Challenge", "6. Backtracking Pattern"),
     ("Recursion + Backtracking", "6. Backtracking Pattern (revision)")],
 7: [("Binary Search vs", "4. Choosing the Right Approach"), ("Binary Search +", "4. Choosing the Right Approach"),
     ("Mixed Revision", "4. Choosing the Right Approach"), ("Choosing", "4. Choosing the Right Approach"),
     ("Binary Search on Answer", "1. Binary Search on Answer Space"), ("Feasibility", "1. Binary Search on Answer Space"),
     ("Greedy", "2. Greedy Algorithms")],
 8: [("Tree Traversal", "2. Tree Traversals"), ("Tree Recursion", "4. Common Tree Problems (height, invert, symmetry, paths)"),
     ("Binary Search Tree", "3. Binary Search Tree"), ("Tree Construction", "3. BST / 4. building a tree from traversals"),
     ("Tree DFS", "4. Common Tree Problems"), ("Path Sum", "4. Common Tree Problems"), ("Lowest", "4. Common Tree Problems — LCA"),
     ("Tree Transformation", "4. Common Tree Problems"), ("Heap", "5. Heaps / Priority Queues"), ("Top-K", "5. Heaps — Top-K"),
     ("K-way", "5. Heaps — merge K sorted"), ("Scheduling", "5. Heaps / Priority Queues")],
 9: [("Representation", "1. Graph Basics"), ("Cycles in Directed", "9. Strongly Connected Components (SCC intuition)"),
     ("Graph Traversal", "2. Graph Traversals"), ("Grid DFS", "2. Graph Traversals"), ("Multi-source", "2. Graph Traversals — multi-source BFS"),
     ("BFS Shortest", "2. Graph Traversals — BFS"), ("State-space", "2. Graph Traversals — BFS (capstone)"),
     ("Grid Components", "2. Graph Traversals (capstone)"), ("Graph Modelling", "2. Graph Traversals (capstone)"),
     ("Degree Reasoning", "3. Cycle Detection (undirected)"), ("Cycle Detection", "3. Cycle Detection"),
     ("Topological", "4. Topological Sort"), ("Bipartite", "5. Bipartite Check"), ("Shortest Path", "6. Shortest Path Algorithms"),
     ("Minimum Spanning", "7. Minimum Spanning Tree"), ("Union Find", "8. Disjoint Set Union")],
 10: [("Greedy vs DP", "3. Greedy vs Dynamic Programming"), ("Dynamic Programming (1D", "1./2. Dynamic Programming (1D)"),
      ("DP", "2. Types of DP Problems"), ("Dynamic Programming", "2. Types of DP Problems"),
      ("Trie", "4. Tries"), ("Bit", "5. Bit Manipulation"), ("XOR", "5. Bit Manipulation — XOR / Single Number"),
      ("String Matching", "6. String Matching — KMP"), ("Rolling Hash", "6. String Matching — Rabin-Karp"),
      ("Design", "(not in the table — see Flags & Gaps)")],
}


def concept_for(belt, pattern):
    for key, topic in RULES[belt]:
        if key in pattern:
            return topic
    raise KeyError((belt, pattern))


def _style(ws, widths, header_row=1, wrap_cols=()):
    head = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=header_row + 1):
        for c in row:
            c.alignment = Alignment(wrap_text=(c.column in wrap_cols), vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def _mix(rows):
    c = collections.Counter(r["difficulty"] for r in rows)
    return c["Easy"], c["Medium"], c["Hard"]


def write(path, master_path, original, out, log, removed, flags, gaps):
    road = {}
    concept_rows = []
    boa = []
    if master_path:
        wb = load_workbook(master_path, read_only=True, data_only=True)
        for r in list(wb["Dojo_Session_Roadmap"].iter_rows(values_only=True))[1:]:
            if r[0] and r[2] and str(r[0]).replace(".", "", 1).isdigit():
                road[(int(float(r[0])), int(float(r[2])))] = (r[3] or "", r[5] or "")
        concept_rows = [list(r) for r in wb["Kalvium Dojo - Concept Mapping"].iter_rows(values_only=True) if any(c is not None for c in r)]
        boa = [r for r in wb["Also On BOA"].iter_rows(values_only=True) if r and r[0] is not None]

    new = Workbook()
    flag_by = collections.defaultdict(list)
    for b, w, n, t in flags:
        flag_by[(b, w, n)].append(t)
    log_by = {(l["belt"], l["week"]): l for l in log}
    out_by_week = collections.defaultdict(list)
    for r in out: out_by_week[(r["belt"], r["week"])].append(r)
    old_by_week = collections.defaultdict(list)
    for r in original: old_by_week[(r["belt"], r["week"])].append(r)

    # ---- README ----------------------------------------------------------------
    ws = new.active; ws.title = "README"
    lines = [
        ["Kalvium Dojo — Group 1 Coding Hours: five practice days per week"],
        [""],
        ["Why", "The weekly concept-session day carries no questions, so each belt-week has 5 practice days x 4 questions = 20, not 6 x 4 = 24."],
        ["Result", f"{len(out)} questions in {len({(r['belt'], r['week'], r['day']) for r in out})} days across 40 weeks (was 960 / 240). {len(removed)} questions were removed; nothing outside the original sheet was added."],
        ["Method", "One decision per belt-week. Exactly one day is dropped and the other five stay as designed (same four questions, same Easy -> Hard order). Which day goes is judged against the Belt Concept Mapping table and the weekly session roadmap: a day outside the belt's listed topics goes first, then self-practice days, then the most redundant day; the challenge day is kept whenever it is in scope."],
        ["Exceptions", "Two weeks (Belt 4 Week 1, Belt 6 Week 1) had two off-concept days, so both were removed and ONE clean day was composed from their in-scope questions. Three single-question swaps replace a clear off-concept question with an in-concept one from a dropped day. All are logged in 'Selection Log'."],
        ["Rules kept", "4 questions per day, ordered Easy -> Medium -> Hard; Belts 1-3 have no Hard problems; no problem appears twice; every problem comes from the original sheet."],
        [""],
        ["Sheets"],
        ["Daily Questions", "The final plan, in the same columns as the original sheet (Day is now 1-5; Usage is renumbered)."],
        ["Selection Log", "Per belt-week: which day was dropped and why, any swaps, and the Easy/Medium/Hard mix before and after."],
        ["Concept Check", "Every kept day mapped to a numbered topic of the Belt Concept Mapping table, with a Review marker where a question needs your eye."],
        ["Flags & Gaps", "Questions that are only borderline for their belt, and table topics the sheet has no (or thin) coverage for."],
        ["Removed Questions", "The 160 questions taken out, each with its original position and the reason."],
        ["Weekly Plan / Belt Summary", "Updated counts per week and per belt."],
        ["Also On BOA", "The original BOA-overlap list, filtered to the kept questions, with their new Day."],
        ["Original 6-day (backup)", "The unchanged 960-question sheet, for comparison or rollback."],
        ["Concept Mapping (reference)", "A copy of the concept table the picks were checked against."],
    ]
    for row in lines: ws.append(row)
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 140
    ws["A1"].font = Font(bold=True, size=14)
    for r in ws.iter_rows(min_row=2):
        r[0].font = Font(bold=True)
        for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- Daily Questions -----------------------------------------------------------
    ws = new.create_sheet("Daily Questions")
    ws.append(["LeetCode #", "Problem", "Difficulty", "Topic", "Pattern", "Day Focus", "URL", "Belt", "Week", "Day", "Q#", "Usage"])
    for r in out:
        ws.append([r["leetcodeNumber"], r["title"], r["difficulty"], r["topic"], r["pattern"], r["dayFocus"], r["url"],
                   r["belt"], r["week"], r["day"], r["position"], r["usage"]])
    _style(ws, [11, 46, 11, 24, 36, 60, 60, 6, 6, 6, 5, 26])

    # ---- Selection Log -----------------------------------------------------------------
    ws = new.create_sheet("Selection Log")
    ws.append(["Belt", "Belt name", "Week", "Session topic (roadmap)", "Action", "Original day(s) removed", "Reason",
               "Swaps", "Qs before", "Qs after", "E/M/H before", "E/M/H after", "Left to self-practice (roadmap)"])
    for l in log:
        k = (l["belt"], l["week"])
        swaps = "; ".join(f"D{s['out'][0]}Q{s['out'][1]} <-> D{s['inn'][0]}Q{s['inn'][1]}: {s['why']}" for s in l["swaps"])
        action = "Two days removed, one clean day composed" if l["composed"] else "One day dropped"
        ws.append([l["belt"], BELT_NAMES[l["belt"]], l["week"], road.get(k, ("", ""))[0], action,
                   ", ".join(f"Day {d}" for d in l["droppedDays"]), l["reason"], swaps,
                   len(old_by_week[k]), len(out_by_week[k]), "/".join(map(str, _mix(old_by_week[k]))),
                   "/".join(map(str, _mix(out_by_week[k]))), road.get(k, ("", ""))[1] if road.get(k, ("", ""))[1] not in ("None", None) else ""])
    _style(ws, [6, 14, 6, 44, 24, 14, 90, 60, 8, 8, 12, 12, 36], wrap_cols=(4, 5, 7, 8, 13))

    # ---- Concept Check -----------------------------------------------------------------------
    ws = new.create_sheet("Concept Check")
    ws.append(["Belt", "Belt name", "Week", "Day", "Day focus", "Patterns in the day", "Concept-table topic matched", "Status", "Note"])
    days = collections.defaultdict(list)
    for r in out: days[(r["belt"], r["week"], r["day"])].append(r)
    review = 0
    for (b, w, d) in sorted(days):
        qs = days[(b, w, d)]
        topics = []
        for q in qs:
            t = concept_for(b, q["pattern"])
            if t not in topics: topics.append(t)
        notes = []
        for q in qs:
            for t in flag_by.get((b, w, q["leetcodeNumber"]), []):
                notes.append(f"#{q['leetcodeNumber']} {q['title']}: {t}")
        status = "Review" if notes else "In belt concept list"
        review += 1 if notes else 0
        ws.append([b, BELT_NAMES[b], w, d, qs[0]["dayFocus"], "; ".join(dict.fromkeys(q["pattern"] for q in qs)),
                   " | ".join(topics), status, "\n".join(notes)])
    _style(ws, [6, 14, 6, 5, 52, 52, 58, 20, 90], wrap_cols=(5, 6, 7, 9))
    for row in ws.iter_rows(min_row=2):
        if row[7].value == "Review":
            row[7].fill = PatternFill("solid", fgColor="FFE699")

    # ---- Flags & Gaps -----------------------------------------------------------------------------
    ws = new.create_sheet("Flags & Gaps")
    ws.append(["Type", "Belt", "Week", "LeetCode #", "Problem / Concept-table topic", "Note"])
    titles = {r["leetcodeNumber"]: r["title"] for r in original}
    for b, w, n, t in flags:
        ws.append(["Borderline question", b, w, n, titles.get(n, ""), t])
    for b, topic, note in gaps:
        ws.append(["Topic with thin/no coverage", b, "", "", topic, note])
    _style(ws, [26, 6, 6, 11, 62, 110], wrap_cols=(5, 6))

    # ---- Removed Questions -------------------------------------------------------------------------------
    ws = new.create_sheet("Removed Questions")
    ws.append(["Belt", "Week", "Original Day", "Q#", "LeetCode #", "Problem", "Difficulty", "Why removed"])
    for r in sorted(removed, key=lambda r: (r["belt"], r["week"], r["day"], r["position"])):
        ws.append([r["belt"], r["week"], r["day"], r["position"], r["leetcodeNumber"], r["title"], r["difficulty"], r["reason"]])
    _style(ws, [6, 6, 12, 5, 11, 50, 11, 110], wrap_cols=(8,))

    # ---- Weekly Plan / Belt Summary ---------------------------------------------------------------------------
    ws = new.create_sheet("Weekly Plan")
    ws.append(["Belt", "Belt name", "Week", "Session topic", "Days", "Questions", "Easy", "Medium", "Hard"])
    for (b, w) in sorted(out_by_week):
        e, m, h = _mix(out_by_week[(b, w)])
        ws.append([b, BELT_NAMES[b], w, road.get((b, w), ("", ""))[0], 5, len(out_by_week[(b, w)]), e, m, h])
    _style(ws, [6, 14, 6, 70, 6, 10, 6, 8, 6], wrap_cols=(4,))

    ws = new.create_sheet("Belt Summary")
    ws.append(["Belt", "Belt name", "Weeks", "Questions (was)", "Questions (now)", "Easy (was)", "Medium (was)", "Hard (was)", "Easy (now)", "Medium (now)", "Hard (now)", "% Hard (now)"])
    for b in range(1, 11):
        o = [r for r in original if r["belt"] == b]; n = [r for r in out if r["belt"] == b]
        oe, om, oh = _mix(o); ne, nm, nh = _mix(n)
        ws.append([b, BELT_NAMES[b], len({r["week"] for r in n}), len(o), len(n), oe, om, oh, ne, nm, nh, round(nh / len(n), 3)])
    oe, om, oh = _mix(original); ne, nm, nh = _mix(out)
    ws.append(["Total", "", 40, len(original), len(out), oe, om, oh, ne, nm, nh, round(nh / len(out), 3)])
    _style(ws, [7, 16, 7, 14, 14, 10, 12, 10, 10, 12, 10, 12])

    # ---- Also On BOA ---------------------------------------------------------------------------------------------------
    ws = new.create_sheet("Also On BOA")
    ws.append(["LeetCode #", "Problem", "Difficulty", "Belt (this sheet)", "Week", "Day (now)", "BOA belt", "BOA question title", "BOA slug"])
    where = {r["leetcodeNumber"]: r for r in out}
    for r in boa[1:]:
        n = int(float(r[0]))
        if n in where:
            ws.append([n, r[1], r[2], where[n]["belt"], where[n]["week"], where[n]["day"], r[6], r[7], r[8]])
    _style(ws, [11, 46, 11, 8, 6, 10, 9, 44, 40])

    # ---- Original backup ----------------------------------------------------------------------------------------------------------
    ws = new.create_sheet("Original 6-day (backup)")
    ws.append(["LeetCode #", "Problem", "Difficulty", "Topic", "Pattern", "Day Focus", "URL", "Belt", "Week", "Day", "Q#", "Usage"])
    for r in original:
        ws.append([r["leetcodeNumber"], r["title"], r["difficulty"], r["topic"], r["pattern"], r["dayFocus"], r["url"],
                   r["belt"], r["week"], r["day"], r["position"], r["usage"]])
    _style(ws, [11, 46, 11, 24, 36, 60, 60, 6, 6, 6, 5, 26])

    # ---- Concept table reference ------------------------------------------------------------------------------------------------------------
    if concept_rows:
        ws = new.create_sheet("Concept Mapping (reference)")
        for row in concept_rows: ws.append(["" if c is None else c for c in row])
        for i, w in enumerate([8, 24, 28, 90, 70, 50], start=1): ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows():
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A5"

    new.save(path)
    print(f"workbook written: {path}  (concept check: {sum(1 for _ in days)} days, {review} marked Review)")
